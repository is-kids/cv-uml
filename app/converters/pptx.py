import io
import logging
import zipfile
from pathlib import Path
from typing import Optional, Union

from PIL import Image

logger = logging.getLogger(__name__)

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".emf", ".wmf"}


def extract_pptx_images(pptx_path: Union[str, Path]) -> list[tuple[int, Image.Image]]:
    pptx_path = Path(pptx_path)

    if not pptx_path.exists():
        logger.error(f"PPTX file not found: {pptx_path}")
        return []

    images = []

    try:
        with zipfile.ZipFile(pptx_path, "r") as zf:
            media_files = [
                name for name in zf.namelist()
                if name.startswith("ppt/media/") and
                Path(name).suffix.lower() in IMAGE_EXTENSIONS
            ]

            media_files.sort()

            for i, media_path in enumerate(media_files):
                try:
                    data = zf.read(media_path)
                    img = Image.open(io.BytesIO(data))
                    images.append((i + 1, img.copy()))
                except Exception as e:
                    logger.warning(f"Failed to load image {media_path}: {e}")

    except zipfile.BadZipFile as e:
        logger.error(f"Invalid PPTX file {pptx_path}: {e}")
    except Exception as e:
        logger.error(f"Failed to extract PPTX images: {e}")

    return images


def extract_pptx_with_slides(pptx_path: Union[str, Path]) -> list[tuple[int, list[Image.Image]]]:
    pptx_path = Path(pptx_path)

    if not pptx_path.exists():
        logger.error(f"PPTX file not found: {pptx_path}")
        return []

    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        logger.warning("python-pptx not available, falling back to simple extraction")
        images = extract_pptx_images(pptx_path)
        return [(1, [img for _, img in images])]

    slides_images = []

    try:
        prs = Presentation(pptx_path)

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_imgs = []

            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image_blob = shape.image.blob
                        img = Image.open(io.BytesIO(image_blob))
                        slide_imgs.append(img.copy())
                    except Exception as e:
                        logger.warning(f"Failed to extract image from slide {slide_num}: {e}")

            if slide_imgs:
                slides_images.append((slide_num, slide_imgs))

    except Exception as e:
        logger.error(f"Failed to process PPTX with python-pptx: {e}")
        images = extract_pptx_images(pptx_path)
        return [(1, [img for _, img in images])]

    return slides_images


def get_pptx_slide_count(pptx_path: Union[str, Path]) -> int:
    try:
        from pptx import Presentation
        prs = Presentation(pptx_path)
        return len(prs.slides)
    except Exception:
        return 0
